import streamlit as st
from datetime import datetime, date
import json
import os
import requests
import time
import io

# 嘗試載入 Word (docx) 與 PDF 解析庫
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pypdf
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

# 頁面基本設定
st.set_page_config(page_title="小學數學科校本 AI 輔助平台 (DeepSeek R1 旗艦版)", layout="wide", page_icon="📐")

# 設定 OpenRouter API Key
st.sidebar.header("🔑 系統設定")
openrouter_api_key = (
    st.secrets.get("OPENROUTER_API_KEY") 
    or st.sidebar.text_input("OpenRouter API Key (sk-or-v1-...)", type="password")
)

# 科組老師名單
TEACHERS_LIST = [
    "鄧慧姍", "陳月娥", "朱嘉欣", "韓斐", "黃群英", 
    "梁倩玉", "容佩誼", "鄭棋昌", "李嘉琪", "林莉雅", 
    "陳詩雅", "陳岍柔", "陳愷彤", "廖心妍"
]

# 歷史紀錄檔案路徑
HISTORY_FILE = "meeting_notes_history.json"
GAMES_FILE = "interactive_games.json"

def load_data(file_path):
    if os.path.exists(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            try:
                return json.load(f)
            except:
                return []
    return []

def save_data(file_path, data):
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def clean_ai_response(text):
    """清理 AI 回傳內容頭尾的 Markdown 代碼塊標記及 DeepSeek R1 思考過程 (<think>...</think>)"""
    if not text:
        return ""
    text = text.strip()
    
    # 清理 DeepSeek R1 的推理思考過程區塊
    if "<think>" in text and "</think>" in text:
        text = text.split("</think>")[-1].strip()
        
    if text.startswith("```html"):
        text = text[7:]
    elif text.startswith("```markdown"):
        text = text[11:]
    elif text.startswith("```"):
        text = text[3:]
    if text.endswith("```"):
        text = text[:-3]
    return text.strip()

def run_openrouter_ai(prompt, system_prompt, api_key):
    """透過 OpenRouter 呼叫免費的 DeepSeek R1 大模型"""
    url = "https://openrouter.ai/api/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    
    # 設定為 DeepSeek R1 免費版模型
    model_name = "deepseek/deepseek-r1:free"
    
    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1,
        "max_tokens": 3000
    }
    try:
        response = requests.post(url, headers=headers, json=payload, timeout=90)
        res_json = response.json()
        if "choices" in res_json and len(res_json["choices"]) > 0:
            return True, res_json["choices"][0]["message"]["content"]
        else:
            err_msg = res_json.get("error", {}).get("message", "未知錯誤")
            return False, f"OpenRouter 回應失敗：{err_msg}"
    except Exception as e:
        return False, f"連線失敗：{str(e)}"

st.title("📐 小學數學科校本 AI 輔助與教材平台 (DeepSeek R1 旗艦版)")

tab1, tab2, tab3 = st.tabs(["🎙️ 集體備課紀錄生成", "📚 歷年備課紀錄庫", "🔗 課堂互動教材庫 (連結版)"])

# ==========================================
# Tab 1: 集體備課紀錄生成 (DeepSeek R1 驅動)
# ==========================================
with tab1:
    st.header("🎙️ 集體備課會議紀錄生成 (貼上 Gemini 逐字稿 + 備課手冊)")
    
    if not openrouter_api_key:
        st.warning("⚠️ 提示：請在左側邊欄 (Sidebar) 或 Secrets 設定 OPENROUTER_API_KEY。")

    # --- 基本資料選擇區 ---
    st.subheader("📝 會議基本資料設定")
    col_g1, col_g2, col_g3 = st.columns(3)
    with col_g1:
        selected_grade = st.selectbox("📌 請選擇年級：", ["一", "二", "三", "四", "五", "六"], index=4)
    with col_g2:
        selected_school_year = st.selectbox("📅 請選擇學年：", ["2025-2026", "2026-2027"], index=1)
    with col_g3:
        selected_meeting_date = st.date_input("🗓️ 請選擇會議日期：", value=date.today())

    col_t1, col_t2 = st.columns(2)
    with col_t1:
        selected_attendees = st.multiselect(
            "👥 請勾選出席老師：", 
            options=TEACHERS_LIST,
            default=["鄧慧姍", "陳月娥", "容佩誼", "陳愷彤"]
        )
    with col_t2:
        selected_recorder = st.selectbox(
            "✍️ 請選擇紀錄老師：", 
            options=TEACHERS_LIST,
            index=2
        )

    st.divider()

    # --- 輸入區：貼上 Gemini 逐字稿 + 上傳 PDF 備課手冊 ---
    st.subheader("📥 輸入會議資料與備課手冊")
    
    col_input1, col_input2 = st.columns([3, 2])
    
    with col_input1:
        transcript_input = st.text_area(
            "1️⃣ 請貼上 Gemini 轉寫好的廣東話/中文逐字稿：", 
            placeholder="請在此處貼上從 Gemini 複製過來的逐字稿內容...", 
            height=250
        )
        
    with col_input2:
        guide_file = st.file_uploader("2️⃣ 📘 (可選) 上傳校本備課手冊 / 課題指引 (.pdf / .docx)", type=["pdf", "docx", "txt", "md"])

    # 解析上傳的備課手冊內容
    guide_text = ""
    if guide_file:
        guide_ext = guide_file.name.split(".")[-1].lower()
        if guide_ext == "pdf":
            if HAS_PDF:
                try:
                    pdf_reader = pypdf.PdfReader(io.BytesIO(guide_file.getvalue()))
                    page_texts = [page.extract_text() for page in pdf_reader.pages if page.extract_text()]
                    guide_text = "\n".join(page_texts)
                except Exception as e:
                    st.error(f"❌ PDF 解析失敗：{str(e)}")
            else:
                st.error("⚠️ 伺服器缺少 pypdf 模組。")
        elif guide_ext == "docx":
            if HAS_DOCX:
                doc = docx.Document(guide_file)
                guide_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                for table in doc.tables:
                    for row in table.rows:
                        guide_text += "\n" + " | ".join([cell.text.strip() for cell in row.cells])
            else:
                st.error("⚠️ 伺服器缺少 python-docx。")
        else:
            guide_text = guide_file.getvalue().decode("utf-8")
        
        if guide_text:
            st.success(f"📖 已讀取校本備課手冊：「{guide_file.name}」（共 {len(guide_text)} 字元）")

    st.divider()

    if st.button("🚀 開始結合逐字稿與手冊生成校本紀錄", type="primary"):
        if not openrouter_api_key:
            st.error("❌ 請先填寫 OpenRouter API Key！")
        elif not transcript_input.strip():
            st.error("❌ 請先在左側框貼上逐字稿內容！")
        else:
            attendees_str = "、".join(selected_attendees) if selected_attendees else "全體數學科老師"
            recorder_str = selected_recorder
            formatted_date_str = selected_meeting_date.strftime("%d-%m-%Y")

            with st.spinner("🧠 DeepSeek R1 正在深度推理逐字稿邏輯並生成校本紀錄..."):
                clean_transcript = transcript_input[:8000].replace("{", "(").replace("}", ")")
                clean_guide = guide_text[:6000].replace("{", "(").replace("}", ")") if guide_text else "無提供額外手冊，請純粹根據逐字稿詳細內容展開教學程序。"

                sys_prompt = "你是一位專業的小學數學教案專家，擅長根據會議逐字稿整理高質量、具體紮實的條列式教案，絕不輸出重複套話。"
                prompt = f"""
你是一位香港資深小學數學科科主席（CDC 課程專家）。
請閱讀【會議逐字稿】，精準提煉老師們討論的**單一特定課題（例如：三角形的面積）**，撰寫一份高品質、條理分明且極具教導性的「集體備課紀錄」。

【絕對禁止機械式重複！嚴格寫作限制規則】：
1. **結構歸納（最多 3 至 4 行 <tr>）**：
   - 切勿將每個小迷思都單獨拆成一行！請把全篇逐字稿歸納為 2 至 3 個綜合教學重點（例如：1. 觀念與底高對應、2. 面積公式與分割探究、3. 已知面積倒推計算）。
2. **具體且紮實的教學程序（Point Form）**：
   - **絕對禁止每一行都複製貼上相同的套話**！
   - 每一列的教學程序必須**完全針對該欄位的教學難點**寫出具體解決策略。必須包含逐字稿中的具體細節（如：圈出直角符號、使用 iPad 平行四邊形分割探究、倒推計算時先乘以 2 再除以底/高、圖像化展示）。
   - 步驟必須使用 `1.` `2.` `3.` 及 `a.` `b.`，且每一個小點之間必須加 `<br>` 換行！
3. **資料來源**：結合備課手冊及逐字稿，寫出具體參考（如：校本備課手冊、進展工作紙、課本動畫及 iPad 探究）。

【第一順位：會議逐字稿】：
{clean_transcript}

【第二順位：參考校本備課手冊】：
{clean_guide}

【請輸出以下 HTML 表格格式】：
### （ {selected_grade} ）年級數學科備課紀錄({selected_school_year})

**單元：** [AI 自動歸納，如：面積]
**課題：** [AI 自動歸納，如：三角形的面積]
**日期：** {formatted_date_str}  
**出席老師：** {attendees_str}  
**紀錄老師：** {recorder_str}  

<table border='1' style='width:100%; border-collapse:collapse; text-align:left;'>
  <tr style='background-color:#f2f2f2;'>
    <th style='width:30%; padding:8px;'>教學重點 / 難點</th>
    <th style='width:50%; padding:8px;'>教學程序 / 解決方法</th>
    <th style='width:20%; padding:8px;'>資料來源</th>
  </tr>
  <!-- 輸出 2 至 3 列極具針對性、內容充實、含 <br> 換行的 Point Form <tr> 區塊 -->
</table>
"""

                success, raw_out = run_openrouter_ai(prompt, sys_prompt, openrouter_api_key)
                if success:
                    st.session_state["current_note"] = clean_ai_response(raw_out)
                    st.session_state["current_grade"] = selected_grade
                    st.session_state["current_year"] = selected_school_year
                    st.toast("✅ DeepSeek R1 深度紀錄生成成功！", icon="🧠")
                else:
                    st.error(f"❌ 生成失敗：{raw_out}")

    if "current_note" in st.session_state and st.session_state["current_note"]:
        st.divider()
        st.subheader("📋 集體備課紀錄（預覽與對話修訂）")
        
        tab_preview, tab_edit = st.tabs(["👁️ 預覽校本表格", "✏️ 編輯原始碼"])
        
        with tab_edit:
            edited_note = st.text_area("HTML / Markdown 內容編輯區", value=st.session_state["current_note"], height=400)
            st.session_state["current_note"] = edited_note
            
        with tab_preview:
            st.markdown(st.session_state["current_note"], unsafe_allow_html=True)
        
        # --- 打字指令讓 AI 自動修訂表格 ---
        st.subheader("💬 打字指示 AI 自動微調修訂")
        refine_instruction = st.text_input("輸入您希望 AI 修改的指示：", placeholder="例如：請將資料來源改為工作紙 P.12，或把第三點教學程序改得更詳細。")
        
        if st.button("🤖 讓 AI 根據指示重新修訂表格", type="secondary"):
            if refine_instruction.strip():
                with st.spinner("🧠 DeepSeek R1 正在根據您的指示更新表格..."):
                    refine_sys_prompt = "你是一位聽從指令的專業教案修改助理，只輸出修正後的完整教案表格。"
                    refine_prompt = f"""
你是一位香港小學數學教案專家。
請根據【修改指示】，修訂並【完整重新輸出】一份最新的 Markdown / HTML 備課紀錄。

【修改指示】：
{refine_instruction}

【原本的備課紀錄表格內容】：
{st.session_state["current_note"]}

【寫作要求】：
1. 請保持相同的 3 欄 HTML <table> 結構、標題、單元與課題。
2. 嚴格執行修改指示，調整相對應的教學重點、教學程序或資料來源。
3. 教學程序繼續保持 Point Form 清單，小點之間使用 <br> 換行。
4. 直接輸出修訂後的完整內容，嚴禁加上任何 Markdown 程式碼區塊標籤（如 ```html）或額外的引言備註。
"""
                    success, updated_raw = run_openrouter_ai(refine_prompt, refine_sys_prompt, openrouter_api_key)
                    if success:
                        st.session_state["current_note"] = clean_ai_response(updated_raw)
                        st.toast("✅ 已成功根據指示修訂表格！", icon="✨")
                        time.sleep(0.5)
                        st.rerun()
                    else:
                        st.error(f"❌ 微調失敗：{updated_raw}")
            else:
                st.warning("請先輸入修改指示。")

        st.divider()

        col1, col2 = st.columns(2)
        with col1:
            st.download_button(
                "📥 下載備課紀錄 (.md)", 
                data=st.session_state["current_note"], 
                file_name=f"備課紀錄_{selected_meeting_date.strftime('%Y%m%d')}.md", 
                mime="text/markdown"
            )
        with col2:
            if st.button("💾 儲存至歷年紀錄庫", type="primary"):
                history = load_data(HISTORY_FILE)
                history.append({
                    "id": int(time.time()),
                    "year": st.session_state.get("current_year", selected_school_year),
                    "grade": f"{st.session_state.get('current_grade', selected_grade)}年級",
                    "date": selected_meeting_date.strftime("%Y-%m-%d"),
                    "title": f"（{st.session_state.get('current_grade', selected_grade)}）年級備課紀錄 ({st.session_state.get('current_year', selected_school_year)}) - {selected_meeting_date.strftime('%d/%m/%Y')}",
                    "content": st.session_state["current_note"]
                })
                save_data(HISTORY_FILE, history)
                st.success("✅ 已成功儲存！可在「📚 歷年備課紀錄庫」分頁按年度與年級查閱。")

# ==========================================
# Tab 2: 歷年備課紀錄庫
# ==========================================
with tab2:
    st.header("📚 歷年備課紀錄庫")
    
    # --- 上載 Word / MD ---
    with st.expander("📤 上載既有 Word (.docx) 或 Markdown (.md) 備課紀錄至資料庫", expanded=False):
        col_u1, col_u2 = st.columns(2)
        with col_u1:
            up_year = st.selectbox("請選擇歸檔學年：", ["2025-2026", "2026-2027", "2024-2025"], index=1, key="up_year")
        with col_u2:
            up_grade = st.selectbox("請選擇歸檔年級：", ["一年級", "二年級", "三年級", "四年級", "五年級", "六年級"], index=4, key="up_grade")
            
        up_file = st.file_uploader("選擇 Word 或 Markdown 檔案", type=["docx", "md", "txt"], key="up_doc_file")
        
        if st.button("🚀 匯入並存檔", type="primary", key="btn_import_doc"):
            if up_file:
                file_text = ""
                file_ext = up_file.name.split(".")[-1].lower()
                
                if file_ext == "docx":
                    if not HAS_DOCX:
                        st.error("⚠️ 伺服器缺少 python-docx 模組。")
                    else:
                        doc = docx.Document(up_file)
                        full_paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                        file_text = "\n\n".join(full_paragraphs)
                        
                        for table in doc.tables:
                            file_text += "\n\n<table border='1' style='width:100%; border-collapse:collapse; text-align:left;'>"
                            for row in table.rows:
                                file_text += "<tr>"
                                for cell in row.cells:
                                    file_text += f"<td style='padding:8px;'>{cell.text.strip()}</td>"
                                file_text += "</tr>"
                            file_text += "</table>"
                            
                elif file_ext in ["md", "txt"]:
                    file_text = up_file.getvalue().decode("utf-8")
                
                if file_text:
                    history = load_data(HISTORY_FILE)
                    history.append({
                        "id": int(time.time()),
                        "year": up_year,
                        "grade": up_grade,
                        "date": datetime.now().strftime("%Y-%m-%d"),
                        "title": f"（{up_grade}）匯入備課紀錄 ({up_year}) - {up_file.name}",
                        "content": file_text
                    })
                    save_data(HISTORY_FILE, history)
                    st.success(f"✅ 檔案「{up_file.name}」已成功匯入！")
                    time.sleep(1)
                    st.rerun()

    st.divider()

    history = load_data(HISTORY_FILE)
    if not history:
        st.info("目前尚無儲存的備課紀錄。")
    else:
        st.subheader("🔍 篩選與檢索")
        col_f1, col_f2 = st.columns(2)
        
        available_years = ["全部學年"] + sorted(list(set([item.get("year", "未知學年") for item in history])), reverse=True)
        available_grades = ["全部年級"] + sorted(list(set([item.get("grade", "未知年級") for item in history])))
        
        with col_f1:
            filter_year = st.selectbox("📅 依學年篩選：", available_years)
        with col_f2:
            filter_grade = st.selectbox("📌 依年級篩選：", available_grades)
            
        filtered_history = history
        if filter_year != "全部學年":
            filtered_history = [item for item in filtered_history if item.get("year") == filter_year]
        if filter_grade != "全部年級":
            filtered_history = [item for item in filtered_history if item.get("grade") == filter_grade]
            
        st.divider()
        
        if not filtered_history:
            st.warning("⚠️ 沒有符合條件的備課紀錄。")
        else:
            titles_map = {
                f"[{item.get('year', '')}][{item.get('grade', '')}] {item['date']} - {item['title']}": item
                for item in reversed(filtered_history)
            }
            
            selected_title = st.selectbox("請選擇要查閱或管理的備課紀錄：", list(titles_map.keys()))
            selected_item = titles_map[selected_title]
            
            st.markdown(selected_item["content"], unsafe_allow_html=True)
            
            col_d1, col_d2 = st.columns(2)
            with col_d1:
                st.download_button(
                    "📥 下載此紀錄 (.md)", 
                    data=selected_item["content"], 
                    file_name=f"{selected_item['title']}.md", 
                    mime="text/markdown"
                )
            with col_d2:
                if st.button(f"🗑️ 刪除這筆紀錄 ({selected_item['date']})", type="secondary"):
                    updated_history = [h for h in history if h.get("id") != selected_item.get("id")]
                    save_data(HISTORY_FILE, updated_history)
                    st.success("🗑️ 該紀錄已成功刪除！")
                    time.sleep(1)
                    st.rerun()

# ==========================================
# Tab 3: 課堂互動教材庫 (純 Link 連結版)
# ==========================================
with tab3:
    st.header("🔗 課堂互動教材與外部資源連結庫")
    st.write("同工可在此新增 GeoGebra、Wordwall 或各類課堂互動遊戲的網址 (Link)，點擊按鈕即可於新分頁開啟！")
    
    with st.expander("➕ 新增教材/遊戲網址 (Link)", expanded=False):
        link_title = st.text_input("教材/遊戲名稱", placeholder="例如：寶石屋數字對決遊戲 / GeoGebra 平行四邊形切割演示")
        link_url = st.text_input("網址 (URL)", placeholder="例如：[https://example.com/game](https://example.com/game) 或 GeoGebra 連結")
        link_grade = st.selectbox("適用年級", ["小一", "小二", "小三", "小四", "小五", "小六", "全校通用"])
        
        if st.button("🚀 發布教材連結", type="primary"):
            if link_title and link_url:
                if not (link_url.startswith("http://") or link_url.startswith("https://")):
                    link_url = "https://" + link_url
                
                games = load_data(GAMES_FILE)
                games.append({
                    "id": int(time.time()),
                    "title": link_title,
                    "url": link_url,
                    "grade": link_grade,
                    "date": datetime.now().strftime("%Y-%m-%d")
                })
                save_data(GAMES_FILE, games)
                st.success(f"✅ 教材連結「{link_title}」已成功新增！")
                st.rerun()
            else:
                st.warning("請填寫教材名稱與完整網址。")
                
    st.divider()
    
    games = load_data(GAMES_FILE)
    if not games:
        st.info("💡 目前教材庫尚無連結，點擊上方「新增教材/遊戲網址 (Link)」來建立第一個連結吧！")
    else:
        st.subheader("🎯 選擇課堂教材並開啟連結")
        
        grade_filter = st.selectbox("依年級篩選教材：", ["全校通用/全部", "小一", "小二", "小三", "小四", "小五", "小六"])
        
        filtered_games = games
        if grade_filter != "全校通用/全部":
            filtered_games = [g for g in games if g.get("grade") == grade_filter or g.get("grade") == "全校通用"]
            
        if not filtered_games:
            st.warning("沒有此年級的教材連結。")
        else:
            for idx, g in enumerate(reversed(filtered_games)):
                col_m1, col_m2, col_m3 = st.columns([3, 2, 1])
                with col_m1:
                    st.markdown(f"**[{g.get('grade', '通用')}] {g['title']}**")
                    st.caption(f"新增日期：{g.get('date', '')} | 網址：{g.get('url', '')}")
                with col_m2:
                    target_url = g.get('url', '#')
                    st.markdown(
                        f'''<a href="{target_url}" target="_blank" style="
                            display: inline-block;
                            padding: 8px 16px;
                            background-color: #FF4B4B;
                            color: white;
                            text-decoration: none;
                            border-radius: 6px;
                            font-weight: bold;
                        ">🔗 在新分頁開啟教材</a>''', 
                        unsafe_allow_html=True
                    )
                with col_m3:
                    if st.button("🗑️ 刪除", key=f"del_link_{g['id']}"):
                        updated_games = [item for item in games if item.get("id") != g.get("id")]
                        save_data(GAMES_FILE, updated_games)
                        st.success("已刪除該連結！")
                        time.sleep(0.5)
                        st.rerun()
                st.divider()
